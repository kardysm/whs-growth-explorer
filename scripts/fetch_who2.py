#!/usr/bin/env python3
"""Download WHO zscore-expanded XLSX tables (wfa, lhfa, wflh) and inspect structure."""
import pathlib
import re
import subprocess

ROOT = pathlib.Path(__file__).resolve().parents[1]
RAW = ROOT / "research" / "raw"
OUT = ROOT / "research" / "data" / "who_lms"
(OUT / "raw").mkdir(parents=True, exist_ok=True)
UA = "Mozilla/5.0 (X11; Linux x86_64; rv:128.0) Gecko/20100101 Firefox/128.0"


def curl(url, dest, max_time=180):
    p = subprocess.run(["curl", "-sSL", "--max-time", str(max_time), "-A", UA,
                        "-o", str(dest), "-w", "%{http_code}|%{size_download}", url],
                       capture_output=True, text=True)
    parts = (p.stdout or "").strip().split("|")
    return parts[0] if parts else "0", parts[1] if len(parts) > 1 else "0"


found = {}
for page in ["who_page_wfa.html", "who_page_lhfa.html", "who_page_wflh.html"]:
    x = (RAW / page).read_text(errors="ignore")
    links = re.findall(r'href="(https://cdn\.who\.int/[^"]+\.xlsx[^"]*)"', x)
    keep = [l for l in sorted(set(links)) if "zscore" in l.lower() and "weeks" not in l.lower()]
    found[page] = keep
    print("##", page)
    for l in keep:
        print("   ", l)

for page, links in found.items():
    for l in links:
        name = l.split("?")[0].rsplit("/", 1)[-1]
        dest = OUT / "raw" / name
        code, size = curl(l, dest)
        print(name, "->", code, size)
        if code == "200" and dest.stat().st_size > 5000:
            head = dest.read_bytes()[:4]
            print("    magic:", head)
            try:
                import openpyxl
                wb = openpyxl.load_workbook(dest, read_only=True)
                ws = wb.worksheets[0]
                rows = list(ws.iter_rows(min_row=1, max_row=3, values_only=True))
                print("    sheet:", ws.title, "| dims:", ws.max_row, "x", ws.max_column)
                for r in rows:
                    print("    row:", r)
                wb.close()
            except Exception as e:
                print("    xlsx parse err:", e)
